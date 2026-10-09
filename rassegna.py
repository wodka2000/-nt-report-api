"""
rassegna.py — Mini rassegna stampa quotidiana, in PDF, per stampa fisica.

Pipeline separata dal monitoraggio normativo (monitor.py): qui le fonti generaliste e gli
approfondimenti personali NON passano per lo scoring di rilevanza legale di Haiku né
generano bozze LinkedIn — sono solo selezionati/riassunti e impaginati in un PDF.

Struttura del giornale (richiesta da Niccolò il 2026-09-14):
  1. Prima pagina — titolo di apertura, meteo di domani (Roma + eventuale città di
     viaggio) e adempimenti dello studio dei prossimi 7 giorni, letti dalla cartella Drive
     "Aptus Daily" (dal 2026-10-08; prima c'erano i titoli di tutte le sezioni). Nella
     copia pubblica su nt-report.com meteo e adempimenti sono sostituiti da un quadro
     famoso o da un'immagine AI, a giorni alterni
  2. Professionale — energia, giochi, concessioni, tecnologia (~2 pagine)
  3. Interessi — geopolitica/guerra, difesa IT/UE, politica estera, cavidotti sottomarini,
     estrazione dai fondali marini, spazio, scienza (fisica/biologia/materiali) (~3 pagine,
     solo se c'è qualcosa di rilevante)
  4. Attualità — cronaca, politica/economia generale, sport, spettacolo (~1 pagina)
  5. Varie — eventuali suggerimenti per la città di viaggio, una mostra, un'attività coi
     bambini, un libro, un disco, un film/serie da guardare, un piatto/vino/ristorante
     (per ultimo) — ciascuno con un link reale trovato via web
     search, non inventato (~1 pagina). Niente oroscopo (rimosso il 2026-09-15 su
     richiesta esplicita, cruciverba rimosso il 2026-10-08).

Il job in bot.py chiama genera_rassegna_pdf() ogni mattina, poi carica il risultato su
Drive (monitor.py, _upload_rassegna_to_drive) sovrascrivendo un file placeholder già
esistente — mai creando un file nuovo, per non incappare di nuovo in storageQuotaExceeded.
"""

import asyncio
import logging
import random
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

logger = logging.getLogger(__name__)

BASE_DIR = Path(__file__).parent
RASSEGNA_DIR = BASE_DIR / "rassegna"

# ── Fonti ────────────────────────────────────────────────────────────────────────
# NB: il feed storico corriere.it/rss/homepage.xml serve contenuto fermo al 2024 (sezione
# abbandonata) — si usano le sezioni principali verificate attive (verificato 2026-09-14,
# tutte con pubDate di oggi).
_ATTUALITA_FEEDS = [
    ("Corriere — Politica", "https://www.corriere.it/dynamic-feed/rss/section/Politica.xml"),
    ("Corriere — Economia", "https://www.corriere.it/dynamic-feed/rss/section/Economia.xml"),
    ("Corriere — Cronache", "https://www.corriere.it/dynamic-feed/rss/section/Cronache.xml"),
    ("Corriere — Spettacoli", "https://www.corriere.it/dynamic-feed/rss/section/Spettacoli.xml"),
    ("Il Foglio", "https://naxos.ilfoglio.it/api/v5/rss/stories/latest"),
    ("Dagospia", "https://raw.githubusercontent.com/this1-it/dagospia/main/rss.xml"),
]

_TECNOLOGIA_FEEDS = [
    ("Agenda Digitale", "https://www.agendadigitale.eu/feed/"),
    ("Key4biz", "https://www.key4biz.it/feed/"),
]

_INTERESSI_FEEDS = [
    ("Analisi Difesa", "https://www.analisidifesa.it/feed"),
    ("Foreign Policy", "https://foreignpolicy.com/feed/"),
    ("ANSA — Mondo", "https://www.ansa.it/sito/notizie/mondo/mondo_rss.xml"),
    ("gCaptain", "https://www.gcaptain.com/feed/"),
    ("SpaceNews", "https://spacenews.com/feed/"),
    ("ESA", "https://www.esa.int/rssfeed/Our_Activities"),
    ("Phys.org — Fisica", "https://phys.org/rss-feed/physics-news/"),
    ("Phys.org — Biologia", "https://phys.org/rss-feed/biology-news/"),
    ("Phys.org — Materiali", "https://phys.org/rss-feed/chemistry-news/materials-science/"),
]

_INTERESSI_TEMI = """
- scenari internazionali di guerra e conflitti in corso
- settore difesa italiano ed europeo (industria, procurement, dottrina)
- politica estera e relazioni internazionali
- cavidotti sottomarini (posa, sabotaggi, infrastrutture critiche)
- attività estrattiva dai fondali marini (mining sottomarino, risorse minerarie)
- autorizzazioni spaziali e spazio in generale (lanci, satelliti, normativa spaziale)
- scienza: fisica, biologia, scienza dei materiali (scoperte rilevanti, non divulgazione generica)
""".strip()

_GIUSTIZIA_AMMINISTRATIVA_URL = "https://www.giustizia-amministrativa.it/web/guest/echi-d-europa-ufficio-studi"

_SETTORI_LABEL = {"energia": "Energia", "concessioni": "Concessioni demaniali", "giochi": "Giochi"}

_JUNK_TITLE_PATTERNS = re.compile(r"apre una nuova finestra|^(Facebook|X|Twitter|LinkedIn|Instagram|YouTube)\s*:?\s*$", re.IGNORECASE)
_JUNK_TITLE_EXACT = {
    "comunicati stampa", "servizi digitali", "agenda trasparente", "antifrode e laboratori",
    "amministrazione trasparente", "note legali", "privacy policy", "mappa del sito",
    "ministero dell'economia e delle finanze", "ministero dell'economia e delle finanze mef",
    "contatti e assistenza", "relazioni con gli operatori",
}


def _is_junk_title(title: str) -> bool:
    """Filtra i link di navigazione/menu/social (alcune pagine PA statiche restituiscono
    voci di menu come se fossero notizie) — non fanno danno nel monitoraggio principale
    (Haiku li scarta per rilevanza), ma qui vengono mostrati direttamente senza scoring."""
    if not title:
        return True
    if _JUNK_TITLE_PATTERNS.search(title):
        return True
    if title.strip().lower() in _JUNK_TITLE_EXACT:
        return True
    if len(title.strip()) < 20:
        return True
    return False


def _clean_title(title: str) -> str:
    """Alcune pagine HTML concatenano etichetta+data+titolo senza spazi nel testo del link
    (es. 'Comunicato stampa11/09/2026Testo...') perché BeautifulSoup unisce nodi di testo
    adiacenti — qui si inseriscono spazi/trattino attorno alle date per leggibilità."""
    if not title:
        return title
    title = re.sub(r"(?<=[a-zà-ù])(\d{1,2}/\d{1,2}/\d{4})", r" \1", title)
    title = re.sub(r"(\d{1,2}/\d{1,2}/\d{4})(?=[A-ZÀ-Ù])", r"\1 — ", title)
    return title.strip()


def _recent(items: list[dict], hours: int = 30) -> list[dict]:
    """Filtra gli item con `published` (time.struct_time) nelle ultime `hours` ore.
    Se un item non ha data, viene tenuto (meglio includere che perdere silenziosamente)."""
    cutoff = datetime.now(timezone.utc) - timedelta(hours=hours)
    out = []
    for it in items:
        pub = it.get("published")
        if not pub:
            out.append(it)
            continue
        try:
            pub_dt = datetime(*pub[:6], tzinfo=timezone.utc)
        except Exception:
            out.append(it)
            continue
        if pub_dt >= cutoff:
            out.append(it)
    return out


async def _fetch_feed_list(feeds: list[tuple[str, str]], max_per_fonte: int, hours: int = 30,
                            filter_recent: bool = True) -> list[dict]:
    """Scarica una lista di feed RSS (nome, url), aggiunge il nome fonte a ogni item.
    Nessuno scoring di rilevanza — solo raccolta. Se `filter_recent` è False ritorna anche
    item più vecchi (serve al fallback "sezione mai vuota" di _fetch_settori_professionali)."""
    from monitor import _fetch_rss

    out = []
    for nome, url in feeds:
        try:
            items = await _fetch_rss(url)
        except Exception as e:
            logger.warning(f"Rassegna: fetch fallito per {nome}: {e}")
            items = None
        if not items:
            continue
        candidati = _recent(items, hours) if filter_recent else items
        for it in candidati[:max_per_fonte]:
            it["source_name"] = nome
            out.append(it)
    return out


async def _fetch_attualita(max_per_fonte: int = 5) -> list[dict]:
    return await _fetch_feed_list(_ATTUALITA_FEEDS, max_per_fonte)


_SETTORI_CHAT_ID = "rassegna_settori"  # namespace dedicato in seen_docs, separato da "owner"

# GU, AGCM e Corte Costituzionale sono condivise tra più ambiti (TOPICS_CONFIG le include
# in energia/concessioni/giochi contemporaneamente): pubblicano di tutto, non solo notizie
# in tema. Le fonti di settore vere e proprie (ARERA, PV Magazine, Mondo Balneare, ADM,
# Jamma.it...) sono già mono-tema per costruzione e non hanno bisogno di verifica.
_SHARED_GROUPS = {"GU", "AGCM"}
# "Ministro Protezione Civile e Politiche del Mare — Notizie" pubblica anche notizie di
# Protezione Civile (terremoti, emergenze) insieme a quelle sul mare: va filtrata per
# pertinenza come le altre fonti condivise, non trattata come mono-tema (richiesta di
# Niccolò, 2026-09-15: "non consultazioni ma news, e solo quelle importanti").
_SHARED_NAMES = {"Corte Costituzionale", "Ministro Protezione Civile e Politiche del Mare — Notizie"}


def _is_shared_source(s: dict) -> bool:
    return s.get("group") in _SHARED_GROUPS or s["name"] in _SHARED_NAMES


async def _fetch_source_items(s: dict) -> list[dict]:
    from monitor import _fetch_rss, _fetch_html_links

    try:
        fetched = (
            await _fetch_rss(s["url"], timeout=s.get("timeout", 20))
            if s["type"] == "rss"
            else await _fetch_html_links(s["url"], link_filter=s.get("link_filter"),
                                          timeout=s.get("timeout", 15),
                                          no_verify=s.get("no_verify", False))
        )
    except Exception as e:
        logger.warning(f"Rassegna settori: fetch fallito per {s['name']}: {e}")
        return []
    if not fetched:
        return []
    for it in fetched:
        it["source_name"] = s["name"]
    return fetched


# Soglia di pertinenza per le fonti condivise (GU/AGCM/Corte Costituzionale) nella rassegna
# stampa: più alta della soglia 5/10 usata da monitor.py per le bozze LinkedIn. Lì lo scopo è
# "vale la pena valutarlo per un post", qui è "vale la pena stamparlo la mattina" — una barra
# più esigente, per tenere fuori adempimenti/proroghe/nomine di routine e lasciare solo
# notizie di sistema o aggiornamenti normativi importanti (richiesta di Niccolò, 2026-09-14).
_RASSEGNA_SCORE_MIN = 8


async def _score_and_bucket(items: list[dict], valid_settori: set[str]) -> dict[str, list[dict]]:
    """Le fonti condivise (GU, AGCM, Corte Costituzionale) pubblicano di tutto, non solo
    notizie in tema: ogni item viene classificato con lo stesso scoring Haiku
    (_assess_relevance) già usato dal monitoraggio principale, ma con soglia più alta
    (_RASSEGNA_SCORE_MIN) perché qui il bar è "notizia di sistema", non "spunto per un post"
    — e smistato nel settore assegnato dal modello, non in quello della fonte. Una sola
    valutazione per item, non una per ambito: le fonti condivise vengono fetchate/valutate
    una volta sola a monte."""
    if not items:
        return {}
    from monitor import _assess_relevance

    loop = asyncio.get_event_loop()

    async def _score(it: dict):
        try:
            res = await loop.run_in_executor(None, _assess_relevance, it.get("title", ""), it.get("summary", ""))
        except Exception as e:
            logger.warning(f"Rassegna: scoring pertinenza fallito per {it.get('url','')}: {e}")
            return None
        if res.get("settore") in valid_settori and res.get("score", 0) >= _RASSEGNA_SCORE_MIN:
            return res["settore"], it
        return None

    risultati = await asyncio.gather(*[_score(it) for it in items])
    out: dict[str, list[dict]] = {}
    for r in risultati:
        if r:
            settore, it = r
            out.setdefault(settore, []).append(it)
    return out


async def _fetch_settori_professionali(max_per_settore: int = 6) -> dict[str, list[dict]]:
    """Notizie del giorno per energia/giochi/concessioni dalle fonti già configurate in
    sources.md (stesso meccanismo del monitoraggio principale, TOPICS_CONFIG, incluse le
    riviste di settore già censite lì — PV Magazine, Quotidiano Energia, Staffetta Online,
    Rinnovabili.it, Jamma.it, GiocoNews, Press Giochi, Agimeg, Mondo Balneare), più
    tecnologia da due fonti dedicate non presenti in sources.md. Dedup separato dal
    monitoraggio principale (namespace 'rassegna_settori') per non interferire con le
    bozze LinkedIn, che restano basate sul proprio tracking 'owner'.

    Questa è la sezione professionale di Niccolò: non deve mai restare vuota di notizie
    PERTINENTI (meglio vuota che piena di rumore) — se dopo il filtro di pertinenza non
    resta nulla di recente, si ripiega su cascata (vedi _select_with_fallback) fino a
    mostrare almeno l'ultima notizia pertinente disponibile, anche se non recentissima."""
    from monitor import SOURCES, TOPICS_CONFIG, _get_topic_sources, _is_seen, _mark_seen

    def _select_with_fallback(items: list[dict]) -> list[dict]:
        candidati = [it for it in items if not _is_junk_title(it.get("title", ""))]
        if not candidati:
            return []
        recenti = _recent(candidati)
        freschi = [it for it in recenti if not _is_seen(it["url"], chat_id=_SETTORI_CHAT_ID)]
        for it in freschi:
            _mark_seen(it["url"], it["title"], it.get("source_name", ""), 0, chat_id=_SETTORI_CHAT_ID)
        if freschi:
            return freschi[:max_per_settore]
        # Fallback SOLO su item con una data verificabile (RSS con published_parsed): le
        # pagine ADM sono scaricate via HTML e non hanno una data reale, quindi _recent()
        # le lascia passare sempre per costruzione — senza questo controllo, quando una
        # fonte ADM non pubblica nulla di nuovo (il caso più frequente), il fallback
        # ripescava avvisi anche di un anno prima presentandoli come notizia del giorno
        # (segnalato da Niccolò, 2026-09-15). Meglio una sezione vuota che notizie del 2025
        # spacciate per attuali.
        con_data = [it for it in candidati if it.get("published")]
        return con_data[:max_per_settore]

    topic_keys = ("energia", "concessioni", "giochi")
    label_by_topic = {"energia": "Energia", "concessioni": "Concessioni demaniali", "giochi": "Giochi"}

    # Le fonti condivise tra più ambiti vengono raccolte una sola volta (non una per ambito)
    # per non moltiplicare fetch e chiamate di scoring.
    shared_sources: dict[str, dict] = {}
    trusted_by_topic: dict[str, list[dict]] = {k: [] for k in topic_keys}

    for topic_key in topic_keys:
        # Le pagine "Gioco Distanza" (Normativa e Comunicati) condividono lo stesso menu di
        # navigazione statico del sito ADM: senza un link_filter dedicato in sources.md,
        # _fetch_html_links restituisce le voci di menu invece di notizie datate.
        sources = [s for s in _get_topic_sources(topic_key, SOURCES) if "gioco distanza" not in s["name"].lower()]
        for s in sources:
            if _is_shared_source(s):
                shared_sources[s["name"]] = s
            else:
                trusted_by_topic[topic_key].extend(await _fetch_source_items(s))

    shared_items: list[dict] = []
    for s in shared_sources.values():
        shared_items.extend(await _fetch_source_items(s))

    target_settori = {TOPICS_CONFIG[k]["settori"][0] for k in topic_keys}
    verified_by_settore = await _score_and_bucket(shared_items, target_settori)

    result: dict[str, list[dict]] = {}
    for topic_key in topic_keys:
        target_settore = TOPICS_CONFIG[topic_key]["settori"][0]
        combinati = trusted_by_topic[topic_key] + verified_by_settore.get(target_settore, [])
        result[label_by_topic[topic_key]] = _select_with_fallback(combinati)

    tech_items = await _fetch_feed_list(_TECNOLOGIA_FEEDS, max_per_fonte=8, filter_recent=False)
    result["Tecnologia"] = _select_with_fallback(tech_items)

    return result


_GA_CHAT_ID = "rassegna_ga"


async def _fetch_giustizia_amministrativa() -> list[dict]:
    """Bollettini periodici (es. 'Echi d'Europa') dalla pagina Ufficio Studi — HTML statico
    con link a PDF. La pagina elenca sempre TUTTI i bollettini pubblicati (aggiornata
    raramente): senza dedup comparirebbero identici ogni giorno. Si tiene traccia dei PDF
    già visti e si mostrano solo quelli nuovi rispetto all'ultima esecuzione."""
    from monitor import _fetch_html_links, _is_seen, _mark_seen

    try:
        items = await _fetch_html_links(_GIUSTIZIA_AMMINISTRATIVA_URL, link_filter=[".pdf"])
    except Exception as e:
        logger.warning(f"Rassegna: fetch Giustizia Amministrativa fallito: {e}")
        return []
    if not items:
        return []
    nuovi = [it for it in items if not _is_seen(it["url"], chat_id=_GA_CHAT_ID)]
    for it in nuovi:
        it["source_name"] = "Ufficio Studi — Giustizia Amministrativa"
        _mark_seen(it["url"], it["title"], "Giustizia Amministrativa", 0, chat_id=_GA_CHAT_ID)
    return nuovi[:5]


async def _fetch_interessi() -> list[dict]:
    """Raccoglie un pool ampio da fonti di geopolitica/difesa/spazio/scienza, poi chiede a
    Claude di selezionare solo ciò che è realmente pertinente ai temi di interesse
    (`_INTERESSI_TEMI`) — non è scoring binario come per il monitoraggio legale, è
    selezione editoriale: se non c'è nulla di interessante oggi, la lista può restare
    vuota e la sezione lo dice esplicitamente, senza riempitivo forzato."""
    from bot import call_claude

    candidati = await _fetch_feed_list(_INTERESSI_FEEDS, max_per_fonte=8, hours=36)
    if not candidati:
        return []

    elenco = "\n".join(
        f"{i}. [{it['source_name']}] {it['title']}" for i, it in enumerate(candidati)
    )
    prompt = f"""\
Sei il curatore della sezione "Interessi personali" di una rassegna stampa quotidiana per
un avvocato che segue, per interesse personale (non professionale), questi temi:
{_INTERESSI_TEMI}

Di seguito un elenco di titoli raccolti oggi da varie fonti. Scegli SOLO quelli
realmente pertinenti e di qualità (niente gossip, niente cronaca sportiva, niente
articoli genericamente scientifici da rotocalco) — massimo 16, anche meno se non c'è
abbastanza materiale valido. Per ciascuno indica il tema di riferimento tra:
guerra/geopolitica, difesa, politica estera, cavidotti sottomarini, estrazione fondali
marini, spazio, scienza.

Elenco:
{elenco}

Rispondi SOLO con un array JSON, senza testo aggiuntivo, con oggetti
{{"indice": <int>, "tema": "<uno dei temi sopra>"}}. Se nulla è pertinente, rispondi con [].
"""
    try:
        msg = call_claude(
            model="claude-haiku-4-5",
            max_tokens=1200,
            messages=[{"role": "user", "content": prompt}],
        )
        raw = msg.content[0].text.strip()
        raw = re.sub(r"^```(?:json)?|```$", "", raw.strip(), flags=re.MULTILINE).strip()
        import json
        selezione = json.loads(raw)
    except Exception as e:
        logger.warning(f"Rassegna interessi: selezione Claude fallita: {e}")
        return []

    out = []
    for sel in selezione:
        idx = sel.get("indice")
        if idx is None or not (0 <= idx < len(candidati)):
            continue
        it = dict(candidati[idx])
        it["tema"] = sel.get("tema", "")
        out.append(it)
    return out


def _estrai_json(testo: str):
    """Primo valore JSON (array o oggetto) nella risposta del modello, ignorando i
    blocchi ``` e qualsiasi frase prima o dopo: capita che Claude aggiunga una riga di
    commento, e un json.loads sull'intera risposta fallisce."""
    import json

    decoder = json.JSONDecoder()
    for i, c in enumerate(testo):
        if c in "[{":
            try:
                return decoder.raw_decode(testo, i)[0]
            except json.JSONDecodeError:
                continue
    raise ValueError(f"nessun JSON nella risposta: {testo[:200]!r}")


# Adempimenti in prima pagina (richiesta di Niccolò, 2026-10-08): non più il Google Doc
# "ToDo" della cartella rassegna, ma le due fonti operative dello studio:
#  - il PDF "scadenze dal X al Y" esportato ogni settimana dal gestionale nella cartella
#    Drive "Aptus Daily" (il nome cambia ogni volta: si prende il più recente);
#  - la checklist mobile "ToDo_checklist_mobile" (foglio "Task aperti"), ovunque si trovi
#    tra i file condivisi con il Service Account.
# Solo adempimenti ravvicinati: da domani a 7 giorni, più i task della checklist ancora
# aperti con data passata da non più di 7 giorni (segnalati come scaduti).
_APTUS_FOLDER_ID = "1lEVZGLkr5JyNvpKaVu1rZI_6EBb3YshQ"
_CHECKLIST_NAME = "ToDo_checklist_mobile"
_TODO_GIORNI = 7
# Nomi con cui Niccolò e la segreteria compaiono tra gli assegnatari: in stampa non
# servono (è la sua ToDo), restano solo i collaboratori.
_ASSEGNATARI_DA_OMETTERE = ("travia", "segreteria", "scadenze", "contabile", "studio",
                            "non specificato")


def _scarica_da_drive(drive, file: dict) -> bytes:
    """Scarica un file da Drive; i formati nativi Google vanno esportati, non scaricati."""
    if file["mimeType"] == "application/vnd.google-apps.spreadsheet":
        return drive.files().export(
            fileId=file["id"],
            mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ).execute()
    return drive.files().get_media(fileId=file["id"]).execute()


def _solo_collaboratori(assegnatari: str) -> str:
    nomi = [n.strip() for n in re.split(r"[,/;]", assegnatari or "") if n.strip()]
    nomi = [n for n in nomi if n.lower() != "altro"
            and not any(x in n.lower() for x in _ASSEGNATARI_DA_OMETTERE)]
    return _nomi_unici(" ".join(n.split()) for n in nomi)


def _nomi_unici(nomi) -> str:
    """Il gestionale scrive "Morelli Michele", la checklist "Michele Morelli": stessa
    persona se le parole coincidono a prescindere dall'ordine. Resta la prima grafia."""
    visti = {}
    for n in nomi:
        visti.setdefault(frozenset(n.lower().split()), n)
    return ", ".join(visti.values())


def _fetch_scadenze_studio(drive, da, a) -> list[dict]:
    """Legge le scadenze 'da fare' tra le date da e a (incluse) dal PDF più recente del
    gestionale. Il PDF è una tabella larga con celle su più righe: l'estrazione testuale
    di pypdf mescola le colonne, quindi lo si passa direttamente a Claude come documento
    e si chiede un JSON. Lo stesso adempimento compare spesso due volte (riga di agenda e
    riga di scadenza della stessa pratica): la deduplica la fa il modello."""
    import base64
    from bot import call_claude

    res = drive.files().list(
        q=(f"'{_APTUS_FOLDER_ID}' in parents and mimeType='application/pdf' "
           f"and name contains 'scadenze' and trashed=false"),
        orderBy="modifiedTime desc", fields="files(id,name,mimeType)", pageSize=1,
    ).execute()
    files = res.get("files", [])
    if not files:
        logger.warning(f"Rassegna: nessun PDF 'scadenze' nella cartella {_APTUS_FOLDER_ID}")
        return []
    pdf = _scarica_da_drive(drive, files[0])

    prompt = f"""\
Questo PDF è l'agenda delle scadenze di uno studio legale. Estrai le righe con data
tra il {da.isoformat()} e il {a.isoformat()} (incluse) la cui colonna "Fatto" NON indica
che sono già state fatte.

Aggiungi anche, come voci a sé, le udienze della colonna "Data Prossima Udienza" che
cadono in quello stesso intervallo (testo: "Udienza" + pratica + foro).

Se lo stesso adempimento della stessa pratica nella stessa data compare più volte (es.
una riga con la sola descrizione e una con la descrizione estesa), tienilo UNA volta.

Per ogni adempimento:
- "data": AAAA-MM-GG
- "testo": codice pratica + cliente in breve, poi l'adempimento, in massimo 15 parole
  (es. "B061 Gridspertise — richiesta di chiarimenti gara IRETI")
- "assegnatari": la colonna Assegnatari così com'è
- "trasferta": la città in cui bisogna andare di persona, SOLO se la voce è un'udienza,
  un appuntamento o un sopralluogo in una città diversa da Roma (la città del foro o
  quella indicata nella descrizione); null per depositi, memorie, note scritte, esami di
  pratica e tutto ciò che si fa dallo studio, e null per tutto ciò che è a Roma.

Rispondi SOLO con un array JSON:
[{{"data": "...", "testo": "...", "assegnatari": "...", "trasferta": null}}]"""
    try:
        msg = call_claude(
            model="claude-sonnet-4-6", max_tokens=3000,
            messages=[{"role": "user", "content": [
                {"type": "document", "source": {
                    "type": "base64", "media_type": "application/pdf",
                    "data": base64.b64encode(pdf).decode(),
                }},
                {"type": "text", "text": prompt},
            ]}],
        )
        righe = _estrai_json(msg.content[0].text)
    except Exception as e:
        logger.warning(f"Rassegna: lettura scadenze da {files[0]['name']} fallita: {e}")
        return []

    out = []
    for r in righe:
        try:
            d = datetime.strptime(r["data"], "%Y-%m-%d").date()
        except (KeyError, TypeError, ValueError):
            continue
        if da <= d <= a and r.get("testo"):
            out.append({"data": d, "testo": r["testo"],
                        "chi": _solo_collaboratori(r.get("assegnatari", "")),
                        "trasferta": r.get("trasferta") or None})
    return out


def _fetch_checklist_task(drive, da, a) -> list[dict]:
    """Task aperti della checklist mobile con data entro a e non più vecchia di
    _TODO_GIORNI giorni prima di da. I task senza data restano fuori: senza una data non
    si può dire che siano ravvicinati. Le colonne si cercano per intestazione e non per
    posizione, perché il file viene rigenerato da un altro agente e l'ordine può cambiare;
    l'unica eccezione è la casella "completata", la cui cella d'intestazione è a sua volta
    una casella di controllo (False) e non un testo."""
    import io
    import openpyxl

    res = drive.files().list(
        q=f"name contains '{_CHECKLIST_NAME}' and trashed=false",
        orderBy="modifiedTime desc", fields="files(id,name,mimeType)", pageSize=1,
    ).execute()
    files = res.get("files", [])
    if not files:
        logger.warning(f"Rassegna: '{_CHECKLIST_NAME}' non condiviso con il Service Account")
        return []

    wb = openpyxl.load_workbook(io.BytesIO(_scarica_da_drive(drive, files[0])), data_only=True)
    ws = wb["Task aperti"] if "Task aperti" in wb.sheetnames else wb.worksheets[0]
    righe = ws.iter_rows(values_only=True)
    intestazione = [c.strip().lower() if isinstance(c, str) else c for c in next(righe, ())]

    def col(*nomi):
        for i, h in enumerate(intestazione):
            if isinstance(h, str) and any(h.startswith(n) for n in nomi):
                return i
        return None

    i_fatto = next((i for i, h in enumerate(intestazione)
                    if isinstance(h, bool) or (isinstance(h, str) and h.startswith("completat"))), None)
    i_data, i_pratica, i_task = col("data"), col("pratica"), col("task")
    i_stato, i_resp = col("stato"), col("collaboratore")
    if i_data is None or i_task is None:
        logger.warning(f"Rassegna: intestazioni inattese in {files[0]['name']}: {intestazione}")
        return []

    limite_ritardo = da - timedelta(days=_TODO_GIORNI)
    out = []
    for r in righe:
        if i_fatto is not None and r[i_fatto] is True:
            continue
        if i_stato is not None and "completat" in str(r[i_stato] or "").lower():
            continue
        d = r[i_data]
        if not isinstance(d, datetime) or not r[i_task]:
            continue
        d = d.date()
        if not (limite_ritardo <= d <= a):
            continue
        pratica = str(r[i_pratica]).strip() if i_pratica is not None and r[i_pratica] else ""
        testo = f"{pratica} — {r[i_task]}" if pratica else str(r[i_task])
        chi = _solo_collaboratori(str(r[i_resp] or "")) if i_resp is not None else ""
        out.append({"data": d, "testo": testo, "chi": chi})
    return out


def _unisci_e_accorcia(voci: list[dict]) -> list[dict]:
    """Le due fonti si sovrappongono (la stessa scadenza è spesso sia nell'agenda del
    gestionale sia nella checklist, scritta in modo diverso) e i task della checklist sono
    paragrafi interi: senza questo passaggio la ToDo sfora su pagina 2 con metà voci
    doppie (anteprima del 2026-10-09). Claude raggruppa i doppioni e riscrive ogni voce in
    una riga; date e collaboratori li ricompone Python dai dati originali, così il modello
    non può spostare una scadenza. Se un indice manca dalla risposta la voce resta com'era:
    meglio una riga lunga che un adempimento perso."""
    from bot import call_claude

    if len(voci) < 2:
        return voci
    elenco = "\n".join(f"{i}. [{v['data'].isoformat()}] {v['testo']}" for i, v in enumerate(voci))
    prompt = f"""\
Ecco gli adempimenti di uno studio legale per i prossimi giorni, presi da due elenchi
diversi che si sovrappongono.

{elenco}

1. Raggruppa le voci che riguardano lo STESSO adempimento della STESSA pratica (stesso
   codice o stesso cliente, stessa cosa da fare), anche se scritte in modo diverso o con
   date diverse. Voci su adempimenti diversi della stessa pratica restano separate.
2. Per ogni gruppo scrivi UNA riga di massimo 14 parole: codice pratica e cliente in breve,
   poi cosa fare. Niente date, niente nomi dei collaboratori, niente numeri di R.G. se
   c'è già il codice pratica.

Ogni indice deve comparire in esattamente un gruppo.
Rispondi SOLO con un array JSON: [{{"indici": [0, 7], "testo": "..."}}]"""
    try:
        msg = call_claude(model="claude-sonnet-4-6", max_tokens=4000,
                          messages=[{"role": "user", "content": prompt}])
        gruppi = _estrai_json(msg.content[0].text)
    except Exception as e:
        logger.warning(f"Rassegna: unione degli adempimenti fallita, uso l'elenco grezzo: {e}")
        return voci

    out, usati = [], set()
    for g in gruppi:
        indici = [i for i in g.get("indici", []) if isinstance(i, int) and 0 <= i < len(voci)
                  and i not in usati]
        if not indici or not g.get("testo"):
            continue
        usati.update(indici)
        chi = _nomi_unici(n for i in indici for n in voci[i]["chi"].split(", ") if n)
        out.append({"data": min(voci[i]["data"] for i in indici), "testo": g["testo"], "chi": chi})
    out += [v for i, v in enumerate(voci) if i not in usati]
    return out


def _fetch_todo_list() -> list[dict]:
    """Adempimenti ravvicinati da scadenze del gestionale + checklist mobile, ordinati
    per data e non ancora uniti (vedi _unisci_e_accorcia, chiamata dopo
    _rileva_citta_viaggio perché l'unione perde il campo "trasferta"). Ogni voce:
    {"data": date, "testo": str, "chi": str, "trasferta": str | None — solo dal PDF}.
    Non solleva mai: se una fonte manca o Drive non risponde, la prima pagina esce con
    quello che c'è."""
    from monitor import _get_google_creds, _DRIVE_SA_FILE

    if not _DRIVE_SA_FILE.exists():
        return []
    try:
        from googleapiclient.discovery import build
        drive = build("drive", "v3", credentials=_get_google_creds())
    except Exception as e:
        logger.warning(f"Rassegna: connessione a Drive per la ToDo fallita: {e}")
        return []

    da = (datetime.now() + timedelta(days=1)).date()
    a = da + timedelta(days=_TODO_GIORNI - 1)
    voci = []
    for fonte in (_fetch_scadenze_studio, _fetch_checklist_task):
        try:
            voci += fonte(drive, da, a)
        except Exception as e:
            logger.warning(f"Rassegna: {fonte.__name__} fallita: {e}")
    return sorted(voci, key=lambda v: v["data"])


def _get_ft_stampa_queue() -> list[dict]:
    """Legge le segnalazioni manuali (FT/La Stampa) non ancora usate e le marca come usate."""
    from monitor import _db_connect, _init_db

    _init_db()
    con = _db_connect()
    rows = con.execute(
        "SELECT id, testo, url FROM rassegna_queue WHERE usato=0 ORDER BY aggiunto_il"
    ).fetchall()
    if rows:
        ids = [r[0] for r in rows]
        con.executemany("UPDATE rassegna_queue SET usato=1 WHERE id=?", [(i,) for i in ids])
        con.commit()
    con.close()
    return [{"testo": r[1], "url": r[2]} for r in rows]


def _save_linkedin_vetrina(items: list[dict]) -> None:
    """Salva i pick di un giro supervisionato su LinkedIn (Niccolò + Claude presenti,
    MAI un cron automatico — vedi linkedin_kb.md e memoria di progetto). Sovrascrive
    solo le voci di oggi (INSERT preceduto da DELETE della stessa data), così rigenerare
    più volte nello stesso giorno non duplica."""
    from monitor import _db_connect
    con = _db_connect()
    con.execute(
        "CREATE TABLE IF NOT EXISTS linkedin_vetrina (data TEXT, testo TEXT, url TEXT, autore TEXT)"
    )
    oggi = datetime.now().strftime("%Y-%m-%d")
    con.execute("DELETE FROM linkedin_vetrina WHERE data=?", (oggi,))
    for it in items:
        con.execute(
            "INSERT INTO linkedin_vetrina (data, testo, url, autore) VALUES (?, ?, ?, ?)",
            (oggi, it.get("testo"), it.get("url"), it.get("autore")),
        )
    con.commit()
    con.close()


def _load_linkedin_vetrina_recente(giorni: int = 5) -> list[dict]:
    """Ritorna i pick dell'ultimo giro su LinkedIn, solo se abbastanza recente (default
    5 giorni) — altrimenti la rubrica sparisce invece di mostrare contenuto vecchio
    spacciato per "dal mio giro" di oggi. Nessuna navigazione automatica qui: questa
    funzione legge solo quello che una sessione supervisionata ha già salvato."""
    from monitor import _db_connect
    con = _db_connect()
    con.execute(
        "CREATE TABLE IF NOT EXISTS linkedin_vetrina (data TEXT, testo TEXT, url TEXT, autore TEXT)"
    )
    ultima_data = con.execute("SELECT MAX(data) FROM linkedin_vetrina").fetchone()[0]
    if not ultima_data:
        con.close()
        return []
    soglia = (datetime.now() - timedelta(days=giorni)).strftime("%Y-%m-%d")
    if ultima_data < soglia:
        con.close()
        return []
    righe = con.execute(
        "SELECT testo, url, autore FROM linkedin_vetrina WHERE data=?", (ultima_data,)
    ).fetchall()
    con.close()
    return [{"testo": t, "url": u, "autore": a} for t, u, a in righe]


# ── Riassunti (il giornale stampato deve essere autosufficiente, non solo link) ────

def _clean_summary_text(raw: str) -> str:
    """Ripulisce un summary RSS da tag HTML/entità/spazi ripetuti."""
    if not raw:
        return ""
    from bs4 import BeautifulSoup
    text = BeautifulSoup(raw, "html.parser").get_text(separator=" ", strip=True)
    return re.sub(r"\s+", " ", text).strip()


async def _fetch_pdf_text(url: str) -> str:
    """Estrae il testo delle prime pagine di un PDF (avvisi ADM, bollettini Ufficio Studi
    sono spesso PDF, non HTML — senza questo, _fetch_content li tratterebbe come pagine
    HTML e produrrebbe solo rumore binario)."""
    import io
    import httpx
    from pypdf import PdfReader

    try:
        async with httpx.AsyncClient(timeout=20, follow_redirects=True) as client:
            resp = await client.get(url, headers={"User-Agent": "Mozilla/5.0"})
        reader = PdfReader(io.BytesIO(resp.content))
        return "\n".join((page.extract_text() or "") for page in reader.pages[:5])[:4000]
    except Exception as e:
        logger.warning(f"Rassegna: estrazione PDF fallita per {url}: {e}")
        return ""


async def _ensure_material(it: dict) -> None:
    """Garantisce che l'item abbia abbastanza materiale testuale per un riassunto fedele:
    usa il summary RSS se sostanzioso, altrimenti scarica il testo della pagina (o del PDF).
    Se anche quello fallisce, il riassunto verrà generato dal solo titolo (nessun dato
    inventato, solo genericità dichiarata al modello in _summarize_items)."""
    from monitor import _fetch_content

    materiale = _clean_summary_text(it.get("summary") or it.get("content") or "")
    if len(materiale) < 60:
        url = it.get("url", "")
        try:
            testo = await _fetch_pdf_text(url) if ".pdf" in url.lower() else (await _fetch_content(url))[0]
            if len(testo) > len(materiale):
                materiale = testo
        except Exception as e:
            logger.warning(f"Rassegna: fetch contenuto fallito per {url}: {e}")
    it["materiale"] = materiale[:1500]


def _summarize_items(items: list[dict]) -> None:
    """Genera un riassunto fattuale di 2-3 frasi per ciascun item (chiave 'riassunto',
    mutato in place). Una sola chiamata Claude per l'intero batch. Il PDF va stampato e
    deve essere autosufficiente: niente voci con solo titolo e link."""
    if not items:
        return
    from bot import call_claude
    import json

    elenco = "\n\n".join(
        f"{i}. Titolo: {it.get('title','')}\n"
        f"Materiale: {it.get('materiale') or '(nessuno — riassumi in modo generico solo dal titolo)'}"
        for i, it in enumerate(items)
    )
    prompt = f"""\
Per ciascuna delle seguenti notizie scrivi un riassunto fattuale di 2-3 frasi (max 50
parole), in italiano, SOLO sulla base del materiale fornito. Non inventare MAI cifre,
nomi, date o dettagli non presenti nel materiale: se il materiale è insufficiente,
riassumi in modo più generico basandoti solo sul titolo, restando onesto sui limiti
dell'informazione disponibile invece di inventare.

{elenco}

Rispondi SOLO con un array JSON: [{{"indice": <int>, "riassunto": "..."}}, ...]
"""
    try:
        msg = call_claude(
            model="claude-haiku-4-5",
            max_tokens=400 + 120 * len(items),
            messages=[{"role": "user", "content": prompt}],
        )
        raw = msg.content[0].text.strip()
        raw = re.sub(r"^```(?:json)?|```$", "", raw.strip(), flags=re.MULTILINE).strip()
        risultati = json.loads(raw)
        for r in risultati:
            idx = r.get("indice")
            if idx is not None and 0 <= idx < len(items):
                items[idx]["riassunto"] = r.get("riassunto", "")
    except Exception as e:
        logger.warning(f"Rassegna: riassunto batch fallito: {e}")


async def _arricchisci_con_riassunti(*gruppi: list[dict]) -> None:
    """Scarica il materiale mancante (concorrente) e genera i riassunti per ogni gruppo di
    item passato (un gruppo = una sezione, per tenere le chiamate Claude ragionevoli)."""
    tutti = [it for gruppo in gruppi for it in gruppo]
    if not tutti:
        return
    await asyncio.gather(*[_ensure_material(it) for it in tutti])
    for gruppo in gruppi:
        _summarize_items(gruppo)


# ── Ricerca web (meteo, link reali per "Varie") ────────────────────────────────────

def _claude_web_search_json(prompt: str, model: str = "claude-sonnet-4-6", max_tokens: int = 1500):
    """Chiama Claude con il tool di web search nativo (server-side, eseguito automaticamente
    dall'API — nessuna esecuzione locale) e ritorna il JSON dell'ultimo blocco di testo della
    risposta. Usato dove serve un dato verificabile (meteo, link reali) invece di un
    suggerimento generico: la ricerca web riduce il rischio di invenzione, ma non lo azzera —
    i link vanno considerati "trovati da un assistente", non garantiti al 100%."""
    from bot import call_claude
    import json

    for tentativo in range(2):
        try:
            msg = call_claude(
                model=model,
                max_tokens=max_tokens,
                tools=[{"type": "web_search_20250305", "name": "web_search", "max_uses": 5}],
                messages=[{"role": "user", "content": prompt}],
            )
            blocchi_testo = [b.text for b in msg.content if getattr(b, "type", None) == "text"]
            if not blocchi_testo:
                raise ValueError("nessun blocco di testo nella risposta")
            raw = blocchi_testo[-1].strip()
            raw = re.sub(r"^```(?:json)?|```$", "", raw, flags=re.MULTILINE).strip()
            try:
                return json.loads(raw)
            except json.JSONDecodeError:
                # Il modello a volte aggiunge prosa prima/dopo il JSON nonostante
                # l'istruzione di rispondere solo con l'oggetto — ultimo tentativo:
                # estrai la sottostringa tra la prima { e l'ultima }.
                match = re.search(r"\{.*\}", raw, flags=re.DOTALL)
                if not match:
                    raise
                return json.loads(match.group())
        except Exception as e:
            if tentativo == 0:
                logger.info(f"Rassegna: ricerca web Claude fallita al primo tentativo, riprovo: {e}")
                continue
            logger.warning(f"Rassegna: ricerca web Claude fallita (2 tentativi): {e}")
            return None


def _rileva_citta_viaggio(todo: list[dict]) -> str | None:
    """Città in cui Niccolò deve essere di persona DOMANI, o None. Dal 2026-10-08 la si
    ricava dal PDF delle scadenze dello studio (campo "trasferta", che Claude compila
    leggendo il PDF solo per udienze/appuntamenti fuori Roma) e non più da un Google Doc
    di appunti personali — richiesta di Niccolò. Una sola città: il caso d'uso è 'domani
    sono in trasferta', non un itinerario multi-tappa. Solo domani, non i giorni dopo:
    altrimenti meteo e consigli di viaggio comparirebbero con giorni di anticipo."""
    domani = (datetime.now() + timedelta(days=1)).date()
    return next((v["trasferta"] for v in todo if v.get("trasferta") and v["data"] == domani), None)

# Condizioni riconosciute: mappate su un'icona SVG disegnata a mano (vedi _METEO_ICONE) —
# niente immagini scaricate da siti meteo (fragili, spesso sprite/JS, difficili da isolare
# come URL singolo affidabile). _condizione_da_dati_verificati sceglie la condizione tra
# queste chiavi in base a dati reali del DOM, il rendering dell'icona è deterministico.
_METEO_ICONE = {
    "sereno": '<svg viewBox="0 0 24 24" width="20" height="20"><circle cx="12" cy="12" r="5" fill="none" stroke="#111" stroke-width="1.5"/><g stroke="#111" stroke-width="1.5"><line x1="12" y1="1" x2="12" y2="4"/><line x1="12" y1="20" x2="12" y2="23"/><line x1="1" y1="12" x2="4" y2="12"/><line x1="20" y1="12" x2="23" y2="12"/><line x1="4.2" y1="4.2" x2="6.3" y2="6.3"/><line x1="17.7" y1="17.7" x2="19.8" y2="19.8"/><line x1="4.2" y1="19.8" x2="6.3" y2="17.7"/><line x1="17.7" y1="6.3" x2="19.8" y2="4.2"/></g></svg>',
    "nuvoloso": '<svg viewBox="0 0 24 24" width="20" height="20"><path d="M6 17a4 4 0 0 1-.5-7.97A5 5 0 0 1 15 8.5 4 4 0 0 1 18 17H6z" fill="none" stroke="#111" stroke-width="1.5"/></svg>',
    "pioggia": '<svg viewBox="0 0 24 24" width="20" height="20"><path d="M6 13a4 4 0 0 1-.5-7.97A5 5 0 0 1 15 4.5 4 4 0 0 1 18 13H6z" fill="none" stroke="#111" stroke-width="1.5"/><g stroke="#111" stroke-width="1.5"><line x1="8" y1="16" x2="7" y2="20"/><line x1="12" y1="16" x2="11" y2="20"/><line x1="16" y1="16" x2="15" y2="20"/></g></svg>',
    "neve": '<svg viewBox="0 0 24 24" width="20" height="20"><path d="M6 13a4 4 0 0 1-.5-7.97A5 5 0 0 1 15 4.5 4 4 0 0 1 18 13H6z" fill="none" stroke="#111" stroke-width="1.5"/><g stroke="#111" stroke-width="1.3"><line x1="8" y1="16" x2="8" y2="21"/><line x1="6" y1="17.5" x2="10" y2="19.5"/><line x1="10" y1="17.5" x2="6" y2="19.5"/><line x1="16" y1="16" x2="16" y2="21"/><line x1="14" y1="17.5" x2="18" y2="19.5"/><line x1="18" y1="17.5" x2="14" y2="19.5"/></g></svg>',
    "nebbia": '<svg viewBox="0 0 24 24" width="20" height="20"><g stroke="#111" stroke-width="1.5" stroke-linecap="round"><line x1="3" y1="8" x2="21" y2="8"/><line x1="3" y1="12" x2="21" y2="12"/><line x1="6" y1="16" x2="18" y2="16"/><line x1="3" y1="20" x2="21" y2="20"/></g></svg>',
    "caldo": '<svg viewBox="0 0 24 24" width="20" height="20"><rect x="10" y="3" width="4" height="12" rx="2" fill="none" stroke="#111" stroke-width="1.5"/><circle cx="12" cy="18" r="3.5" fill="#111" stroke="#111" stroke-width="1.5"/><line x1="12" y1="8" x2="12" y2="15" stroke="#111" stroke-width="1.5"/></svg>',
    "freddo": '<svg viewBox="0 0 24 24" width="20" height="20"><rect x="10" y="3" width="4" height="12" rx="2" fill="none" stroke="#111" stroke-width="1.5"/><circle cx="12" cy="18" r="3.5" fill="none" stroke="#111" stroke-width="1.5"/><g stroke="#111" stroke-width="1.2"><line x1="17" y1="4" x2="21" y2="4"/><line x1="19" y1="2" x2="19" y2="6"/><line x1="17.6" y1="2.6" x2="20.4" y2="5.4"/><line x1="20.4" y1="2.6" x2="17.6" y2="5.4"/></g></svg>',
}


def _slugify_citta(citta: str) -> str:
    sostituzioni = str.maketrans("àèéìòù", "aeeiou")
    return citta.strip().lower().translate(sostituzioni).replace(" ", "-")


def _condizione_da_dati_verificati(rain_class: str, rain_pct: int, temp_c: int) -> str:
    """Deriva la condizione SOLO da segnali verificabili estratti dal DOM (probabilità di
    pioggia, classe pioggia sì/no, temperatura) — niente mappatura dei codici icona
    'icon-awi-white-NN' di meteoam.it: verificato il 2026-09-15 che quei codici non hanno
    una legenda pubblica reperibile (non nei CSS/JS del sito), e indovinarli avrebbe
    significato mostrare un'icona sbagliata (es. sole quando piove) — peggio che non avere
    l'icona. Risultato: distingue solo pioggia/caldo/freddo/sereno, non nuvoloso/neve/nebbia
    (mai verificabili con i dati disponibili)."""
    if "icon-rain" in rain_class and "no-rain" not in rain_class:
        return "pioggia"
    if rain_pct >= 40:
        return "pioggia"
    if temp_c >= 32:
        return "caldo"
    if temp_c <= 3:
        return "freddo"
    return "sereno"


# La classe CSS dell'icona vento (es. 'd-ne') seleziona un'immagine SVG con una freccia che
# punta nella direzione VERSO CUI soffia il vento, non da dove viene — verificato il
# 2026-09-18 ingrandendo l'SVG di 'd-ne' (freccia verso NE, coda verso SW) e confrontandolo
# con il riquadro principale della stessa pagina, che nello stesso momento indicava "Vento:
# 4 Km/h da 55°" (NE) mentre la griglia oraria aveva classe 'd-sw' per quell'ora — i due
# valori sono opposti di ~180° (55° vs SW≈225°), confermando che la classe è la direzione
# "verso" e va invertita per ottenere la convenzione meteorologica "vento DA" (es. libeccio
# = vento da SW). Prima di questo fix riportavamo sempre l'esatto opposto del vento reale.
_INVERTI_DIREZIONE_VENTO = {
    "N": "S", "S": "N", "E": "W", "W": "E",
    "NE": "SW", "SW": "NE", "NW": "SE", "SE": "NW",
    "NNE": "SSW", "SSW": "NNE", "ENE": "WSW", "WSW": "ENE",
    "ESE": "WNW", "WNW": "ESE", "SSE": "NNW", "NNW": "SSE",
}


def _fetch_meteo_playwright(citta: str) -> dict | None:
    """Legge le previsioni reali da meteoam.it renderizzando la pagina con un browser
    headless (Playwright/Chromium): è una SPA JS, l'HTML statico non contiene i dati
    (verificato 2026-09-15). Estrae dal DOM renderizzato: alba/tramonto/umidità dal pannello
    principale, temperatura/vento/probabilità di pioggia orari da '.weather-info-container'
    (che copre abbondantemente anche il giorno dopo). La direzione del vento è codificata
    nel nome classe CSS dell'icona (es. 'd-e-se' = Est-Sudest) ma va invertita di 180° per
    ottenere la direzione "da" — vedi _INVERTI_DIREZIONE_VENTO."""
    from bs4 import BeautifulSoup
    from playwright.sync_api import sync_playwright

    slug = _slugify_citta(citta)
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            # Il server gira in UTC (Etc/UTC): senza forzare il fuso qui, meteoam.it
            # calcola alba/tramonto/etichette orarie in base al fuso del browser e li
            # mostra sfasati di ~2 ore rispetto all'ora reale di Roma (verificato il
            # 2026-09-17 confrontando con il sito: alba vera 06:51, estratta 04:49).
            context = browser.new_context(timezone_id="Europe/Rome", locale="it-IT")
            page = context.new_page()
            page.goto(f"https://www.meteoam.it/it/meteo-citta/{slug}", wait_until="networkidle", timeout=30000)
            page.wait_for_timeout(1500)
            html = page.content()
            browser.close()
    except Exception as e:
        logger.warning(f"Rassegna meteo: rendering Playwright fallito per {citta}: {e}")
        return None

    soup = BeautifulSoup(html, "html.parser")

    def _valore_parametro(tipo: str) -> str | None:
        icona = soup.select_one(f".meteogram-info-list-parameter-icon.{tipo}")
        if not icona:
            return None
        valore = icona.find_next_sibling("span", class_="meteogram-info-list-parameter-value")
        return valore.get_text(strip=True) if valore else None

    alba = _valore_parametro("sunrise")
    tramonto = _valore_parametro("sunset")
    umidita = _valore_parametro("humidity")
    if not alba:
        logger.warning(f"Rassegna meteo: pannello principale non trovato per {citta} — layout cambiato?")
        return None

    containers = soup.select(".weather-info-container")
    if not containers:
        return None

    try:
        prima_ora = int(containers[0].select_one(".weather-info-date").get_text(strip=True)[:2])
    except Exception:
        return None

    oggi = datetime.now().date()
    orari: dict[tuple, dict] = {}
    for i, c in enumerate(containers):
        data_corrente = oggi + timedelta(days=(prima_ora + i) // 24)
        ora = (prima_ora + i) % 24
        temp_el = c.select_one(".weather-info-temperature")
        rain_el = c.select_one(".weather-rain-probability")
        rain_icon_el = c.select_one(".weather-rain-icon")
        wind_icon_el = c.select_one(".wind-icon")
        wind_val_el = c.select_one(".wind-value")
        if not temp_el:
            continue
        try:
            temp_c = int(re.sub(r"[^\d-]", "", temp_el.get_text(strip=True)))
        except ValueError:
            continue
        rain_pct = int(re.sub(r"\D", "", rain_el.get_text())) if rain_el and rain_el.get_text(strip=True) else 0
        rain_class = " ".join(rain_icon_el.get("class", [])) if rain_icon_el else ""
        direzione = ""
        if wind_icon_el:
            classi = [cl for cl in wind_icon_el.get("class", []) if cl.startswith("d-")]
            if classi:
                direzione_verso = classi[0][2:].upper().replace("-", "")
                direzione = _INVERTI_DIREZIONE_VENTO.get(direzione_verso, direzione_verso)
        velocita = wind_val_el.get_text(strip=True) if wind_val_el else ""
        orari[(data_corrente, ora)] = {
            "condizione": _condizione_da_dati_verificati(rain_class, rain_pct, temp_c),
            "temperatura": f"{temp_c}°C",
            "vento_intensita": f"{velocita} km/h" if velocita else "",
            "vento_direzione": direzione,
        }

    domani = oggi + timedelta(days=1)

    def _fascia(ora: int, con_vento: bool = True) -> dict | None:
        dato = orari.get((domani, ora))
        if not dato:
            return None
        return dato if con_vento else {"condizione": dato["condizione"], "temperatura": dato["temperatura"]}

    mattina, pomeriggio, sera, notte = _fascia(9), _fascia(15), _fascia(20), _fascia(23, con_vento=False)
    if not any((mattina, pomeriggio, sera, notte)):
        logger.warning(f"Rassegna meteo: nessun orario di domani trovato nella finestra scaricata per {citta}")
        return None

    return {
        "citta": citta, "alba": alba, "tramonto": tramonto, "umidita": umidita,
        "mattina": mattina, "pomeriggio": pomeriggio, "sera": sera, "notte": notte,
    }


async def _fetch_meteo_dettagliato(citta: str) -> dict | None:
    """Previsioni di domani per una città: condizione, temperatura e vento per mattina/
    pomeriggio/sera, condizione+temperatura per la notte, alba/tramonto, umidità. Unica
    fonte: rendering diretto di meteoam.it (Aeronautica Militare) via Playwright (dati
    reali, non riassunti da un modello — vedi _fetch_meteo_playwright). Niente fallback su
    altre fonti (richiesta esplicita di Niccolò, 2026-09-17, dopo aver visto dati meteo
    inattendibili): se il rendering fallisce, la sezione meteo semplicemente non compare
    quel giorno — meglio assente che imprecisa/da fonte diversa."""
    loop = asyncio.get_event_loop()
    return await loop.run_in_executor(None, _fetch_meteo_playwright, citta)


async def _fetch_meteo(citta_viaggio: str | None) -> list[dict]:
    """Meteo dettagliato di domani per Roma (sempre) e per l'eventuale città di viaggio già
    individuata a monte nella ToDo list (es. 'treno Milano ore 17') — la rilevazione è
    unica e condivisa con _fetch_varie_viaggio per non duplicare la chiamata Claude."""
    citta = ["Roma"]
    if citta_viaggio and citta_viaggio not in citta:
        citta.append(citta_viaggio)

    risultati = []
    for c in citta:
        dati = await _fetch_meteo_dettagliato(c)
        if dati:
            risultati.append(dati)
    return risultati


def _fetch_varie_viaggio(citta: str | None) -> dict | None:
    """Se è stata individuata una città di viaggio nella ToDo list, cerca sul web cosa fare/
    vedere/mangiare lì (link reali, stesso principio di _genera_varie) — richiesta di
    Niccolò, 2026-09-15: 'ricorda anche questo, oltre al meteo, se vedi che sono in
    viaggio'."""
    if not citta:
        return None
    dati = _claude_web_search_json(
        f"""Cerca sul web spunti REALI e verificabili per un avvocato in viaggio a {citta}
domani: cosa fare, cosa vedere, dove mangiare. Per ciascuno un testo BREVE (max 2 frasi, in
italiano) e un link reale trovato con la ricerca (non inventato).

Rispondi SOLO con un oggetto JSON:
{{"citta": "{citta}",
  "fare": {{"testo": "...", "link": "https://..."}},
  "vedere": {{"testo": "...", "link": "https://..."}},
  "mangiare": {{"testo": "...", "link": "https://..."}}}}
Se per un campo non trovi un link reale affidabile, ometti quella chiave piuttosto che
inventarla.""",
        max_tokens=2000,
    )
    if not dati or not dati.get("citta"):
        return None
    _save_varie_storico(dati, prefisso="viaggio_", citta=dati["citta"])
    return dati


# ── Sezione "Varie" (contenuti generati, non notizie) ─────────────────────────────

def _varie_storico_connect():
    """Connessione con la tabella varie_storico garantita, incluse le colonne link/
    immagine/citta aggiunte il 2026-09-29 per la webapp delle varie (spunte, stelle,
    commenti su nt-report.com/varie.html) — prima si salvava solo il testo."""
    from monitor import _db_connect
    con = _db_connect()
    con.execute(
        "CREATE TABLE IF NOT EXISTS varie_storico (data TEXT, chiave TEXT, testo TEXT, "
        "PRIMARY KEY (data, chiave))"
    )
    colonne = {r[1] for r in con.execute("PRAGMA table_info(varie_storico)")}
    for col in ("link", "immagine", "citta"):
        if col not in colonne:
            con.execute(f"ALTER TABLE varie_storico ADD COLUMN {col} TEXT")
    return con


def _save_varie_storico(dati: dict, prefisso: str = "", citta: str | None = None) -> None:
    """Salva i testi di oggi per poterli escludere dai prossimi giorni (stesso pattern
    del vecchio cruciverba, rimosso il 2026-10-08). Senza questo, il prompt di ricerca web non ha modo
    di sapere cosa ha già proposto e tende a riconvergere sulle stesse opzioni "sicure"
    (es. la stessa mostra permanente, lo stesso classico) — segnalato da Niccolò il
    2026-09-16 ("le varie sono molte uguali a quelle di ieri"). Usato anche per le varie
    di viaggio (prefisso "viaggio_"), che servono solo alla webapp: _load_varie_recenti
    le ignora perché non sono tra le chiavi di label_recenti."""
    con = _varie_storico_connect()
    oggi = datetime.now().strftime("%Y-%m-%d")
    for chiave, dato in (dati or {}).items():
        if not isinstance(dato, dict):
            continue  # es. "citta" nelle varie di viaggio
        testo = dato.get("testo")
        if testo:
            con.execute(
                "INSERT OR REPLACE INTO varie_storico (data, chiave, testo, link, immagine, citta) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (oggi, prefisso + chiave, testo, dato.get("link"), dato.get("immagine"), citta),
            )
    con.commit()
    con.close()


def _load_varie_recenti(giorni: int = 14) -> dict[str, list[str]]:
    con = _varie_storico_connect()
    soglia = (datetime.now() - timedelta(days=giorni)).strftime("%Y-%m-%d")
    righe = con.execute(
        "SELECT chiave, testo FROM varie_storico WHERE data >= ? ORDER BY data DESC", (soglia,)
    ).fetchall()
    con.close()
    recenti: dict[str, list[str]] = {}
    for chiave, testo in righe:
        recenti.setdefault(chiave, []).append(testo)
    return recenti


def _genera_varie() -> dict[str, dict]:
    """Un blocco di spunti non-notiziosi: mostra, attività con bambini, libro, disco,
    film/serie, piatto/vino/ristorante (in quest'ordine — "a tavola" per ultimo su
    richiesta di Niccolò). Ogni spunto è cercato sul web (non inventato "a memoria") e
    viene sempre con un link reale verificabile — per "con i bambini" può essere
    un'immagine invece che un link, se più adatta. La ricerca web riduce ma non azzera il
    rischio che un link non sia perfettamente accurato: vale la stessa cautela di
    _claude_web_search_json. Tiene anche uno storico (_save_varie_storico/
    _load_varie_recenti) per non riproporre le stesse cose — senza, il modello tende a
    riconvergere sulle opzioni più "sicure"/note ogni volta (segnalato da Niccolò il
    2026-09-16: "le varie sono molte uguali a quelle di ieri")."""
    recenti = _load_varie_recenti()
    label_recenti = {
        "mostra": "mostre", "attivita_bambini": "attività con i bambini", "libro": "libri",
        "disco": "dischi", "film_serie": "film/serie", "piatto_vino": "piatti/vini/ristoranti",
    }
    blocco_recenti = ""
    if recenti:
        righe = []
        for chiave, nome in label_recenti.items():
            testi = recenti.get(chiave)
            if testi:
                righe.append(f"- {nome}: " + "; ".join(testi[:10]))
        if righe:
            blocco_recenti = (
                "\n\nGIÀ PROPOSTI negli ultimi giorni (NON ripeterli, scegli qualcosa di "
                "diverso per ciascuna categoria):\n" + "\n".join(righe) + "\n"
            )

    dati = _claude_web_search_json(
        f"""\
Cerca sul web spunti REALI e verificabili per la rubrica "Varie" di una rassegna stampa
personale per un avvocato italiano appassionato di cultura generale. Servono 6 elementi,
ciascuno con un testo BREVE (max 2 frasi, in italiano, tono colloquiale ma non sciatto) e
un link reale trovato con la ricerca (non inventato):
{blocco_recenti}
Per mostra, attivita_bambini e piatto_vino (se un ristorante): PREDILIGI opzioni a Roma o
nel Lazio, dove vive l'avvocato — vanno bene opzioni altrove SOLO se eccezionalmente
importanti/note (es. una mostra internazionale di grande rilievo), altrimenti scegli sempre
qualcosa di raggiungibile senza viaggiare.

1. mostra: una mostra/esposizione attualmente visitabile a Roma o nel Lazio (o una
   collezione permanente nota lì), con link al sito ufficiale della mostra/museo
2. attivita_bambini: un'attività o risorsa per bambini, preferibilmente a Roma/Lazio se è
   un'attività fisica (gioco, laboratorio, sito educativo) — se è un libro illustrato o
   un'idea generica la localizzazione non si applica; con link — se più adatta di un link,
   indica anche (o in alternativa) l'URL di un'immagine rappresentativa reale nel campo
   "immagine"
3. libro: un libro (anche un classico) pertinente a diritto, economia, storia o scienza,
   con link a una pagina reale (editore, libreria online, Wikipedia)
4. disco: un disco/album da riascoltare, con link a una pagina reale (Wikipedia, servizio
   di streaming, etichetta)
5. film_serie: un film o una serie tv (al cinema ora, oppure in streaming/casa), con link
   a una pagina reale (sito del cinema/programmazione, piattaforma di streaming, Wikipedia)
6. piatto_vino: una ricetta, una bottiglia o un ristorante reale — se un ristorante,
   preferibilmente a Roma/Lazio — con link alla fonte trovata (sito della ricetta, cantina,
   ristorante)

Rispondi SOLO con un oggetto JSON valido, in questo ordine di chiavi:
{{"mostra": {{"testo": "...", "link": "https://..."}},
 "attivita_bambini": {{"testo": "...", "link": "https://...", "immagine": "https://..." }},
 "libro": {{"testo": "...", "link": "https://..."}},
 "disco": {{"testo": "...", "link": "https://..."}},
 "film_serie": {{"testo": "...", "link": "https://..."}},
 "piatto_vino": {{"testo": "...", "link": "https://..."}}}}
Se per un campo non trovi un link reale affidabile, ometti quella chiave (link o immagine)
piuttosto che inventarla.""",
        model="claude-sonnet-4-6",
        max_tokens=2000,
    )
    dati = dati or {}
    if dati:
        _save_varie_storico(dati)
    return dati


# ── Composizione HTML ──────────────────────────────────────────────────────────────

_CSS = """
  @page { size: A4; margin: 1.6cm; }
  body { font-family: Georgia, 'Times New Roman', serif; color: #111; font-size: 12px; }
  .page-front { page-break-after: always; }
  h1.masthead { font-size: 40px; text-align: center; border-top: 4px solid #111;
                border-bottom: 4px solid #111; padding: 10px 0; margin-bottom: 2px;
                letter-spacing: 2px; font-variant: small-caps; }
  .data { text-align: center; font-style: italic; color: #555; margin-bottom: 16px;
          border-bottom: 1px solid #ccc; padding-bottom: 10px; }
  h2.sezione { font-size: 22px; margin: 0 0 10px 0; border-bottom: 3px solid #111;
               padding-bottom: 4px; font-variant: small-caps; page-break-after: avoid;
               page-break-inside: avoid; }
  h3 { font-size: 13.5px; margin: 0 0 4px 0; text-transform: uppercase;
       letter-spacing: 0.4px; border-bottom: 1px solid #111; padding-bottom: 2px;
       break-after: avoid; break-inside: avoid; }
  a { color: #111; text-decoration: none; font-weight: bold; }
  .fonte { color: #777; font-size: 10px; font-style: italic; }

  /* Un unico flusso continuo a colonne per tutte le sezioni interne (non un blocco per
     sezione): quando una lista finisce a metà pagina, la sezione dopo continua a
     riempire lo spazio invece di saltare a una colonna/pagina nuova — come le pagine
     interne di un giornale vero. CSS `columns` nativo (non split Python: quel primo
     tentativo bilanciava per numero di voci, non per spazio reale, e produceva colonne
     con enormi vuoti quando una voce era più corta delle altre — vedi git history
     2026-09-15). Niente column-span qui: varie e meteo (il meteo lo richiederebbe per
     occupare tutta la larghezza) stanno apposta fuori da questo contenitore, le varie
     sull'ultima pagina (_build_ultima_pagina) e il meteo in prima pagina — quella
     combinazione mandava in crash WeasyPrint in modo intermittente, vedi nota nel codice
     Python del 2026-09-16. */
  .corpo-continuo { columns: 3; column-gap: 22px; column-rule: 1px solid #bbb;
                     text-align: justify; hyphens: auto; orphans: 3; widows: 3;
                     column-fill: auto; }
  .corpo-continuo ul { list-style: none; margin: 0 0 14px 0; padding: 0; }
  .corpo-continuo li { margin: 0; padding: 7px 0; line-height: 1.35; border-top: 1px solid #ddd;
                      break-inside: avoid; }
  .corpo-continuo li:first-child { border-top: none; }
  .corpo-continuo li a { display: block; font-size: 11.5px; margin-bottom: 3px; }
  .corpo-continuo li .riassunto { display: block; font-size: 11px; font-weight: normal;
                                  color: #333; margin-bottom: 3px; }
  .corpo-continuo .fonte { display: block; margin-top: 2px; }

  /* Prima pagina: titolo di apertura a piena larghezza, poi meteo e adempimenti */
  .lead { border-bottom: 2px solid #111; padding-bottom: 14px; margin-bottom: 16px; }
  .lead .tag { font-size: 11px; text-transform: uppercase; color: #888; letter-spacing: 0.5px; }
  .lead .titolo { display: block; font-size: 26px; font-weight: bold; line-height: 1.2;
                  margin-top: 4px; font-family: Georgia, serif; }
  /* Una colonna, una riga per adempimento: con ~30 voci le due colonne CSS non si
     spezzano bene tra le pagine (Chromium sposta tutta la lista sul foglio dopo). */
  .todo { border: 1px solid #111; padding: 8px 14px; margin-bottom: 16px; }
  .todo h3 { border: none; margin-bottom: 4px; }
  .todo ul { list-style: none; margin: 0; padding: 0; }
  .todo li { font-size: 10.5px; line-height: 1.25; padding: 1px 0; border-top: 1px solid #eee;
             break-inside: avoid; }
  .todo li:first-child { border-top: none; }
  .todo .checkbox { font-size: 12px; margin-right: 4px; }
  .todo .todo-data { font-weight: bold; display: inline-block; min-width: 68px; }
  .todo .ritardo { color: #a00; }
  .todo .todo-chi { color: #666; font-style: italic; }
  .arte { text-align: center; margin-top: 10px; }
  .arte img { max-width: 100%; max-height: 15cm; }
  .arte .didascalia { font-size: 10.5px; color: #555; font-style: italic; margin-top: 6px; }

  .varie-block { margin-bottom: 16px; }
  .varie-block .label { font-weight: bold; }
  .varie-block a { font-weight: normal; font-size: 10px; color: #555; }
  .varie-img { max-width: 100%; max-height: 160px; margin-top: 6px; }
  .meteo-blocco .label { display: block; margin-bottom: 4px; }
  .meteo-sole { font-size: 10.5px; color: #555; margin-bottom: 8px; }
  table.meteo-tabella { border-collapse: collapse; width: 100%; }
  td.meteo-fascia { border: 1px solid #ccc; text-align: center; padding: 6px 4px;
                     vertical-align: top; width: 25%; }
  .meteo-fascia-nome { font-size: 10px; text-transform: uppercase; letter-spacing: 0.4px;
                        color: #777; margin-bottom: 4px; }
  .meteo-icona svg { display: block; margin: 0 auto; }
  .meteo-temp { font-weight: bold; font-size: 13px; margin-top: 4px; }
  .meteo-vento { font-size: 9px; color: #555; margin-top: 3px; }
"""


# Il blocco privato della prima pagina (meteo + adempimenti dello studio) è delimitato da
# questi due commenti HTML: _versione_pubblica_pdf sostituisce tutto ciò che sta in mezzo
# con l'immagine del giorno. Commenti e non un <div> da cercare con una regex perché il
# blocco contiene div annidati, su cui un ".*?</div>" si fermerebbe troppo presto.
_PRIVATO_INIZIO = "<!--PRIVATO-INIZIO-->"
_PRIVATO_FINE = "<!--PRIVATO-FINE-->"

_GIORNI_BREVI = ["lun", "mar", "mer", "gio", "ven", "sab", "dom"]


def _build_todo_checklist(todo: list[dict]) -> str:
    if not todo:
        return ""
    domani = (datetime.now() + timedelta(days=1)).date()
    righe = []
    for voce in todo:
        d = voce["data"]
        if d < domani:
            quando = f'<span class="todo-data ritardo">scad. {d.strftime("%d/%m")}</span>'
        else:
            quando = f'<span class="todo-data">{_GIORNI_BREVI[d.weekday()]} {d.strftime("%d/%m")}</span>'
        chi = f' <span class="todo-chi">({voce["chi"]})</span>' if voce.get("chi") else ""
        righe.append(f'<li><span class="checkbox">☐</span> {quando} {voce["testo"]}{chi}</li>')
    return f'<div class="todo"><h3>Da fare</h3><ul>{"".join(righe)}</ul></div>'


def _build_front_page(settori: dict, interessi: list, attualita: list, meteo: list[dict],
                      todo: list[dict]) -> str:
    """Prima pagina: testata, un solo titolo di apertura, poi meteo di domani e
    adempimenti dello studio. Le tre colonne di titoli sono state tolte su richiesta di
    Niccolò (2026-10-08): la prima pagina serve a organizzare la giornata, le notizie
    stanno nelle pagine interne."""
    candidati = (
        [(it, "Professionale") for items in settori.values() for it in items][:1]
        + [(it, "Interessi") for it in interessi[:1]]
        + [(it, "Attualità") for it in attualita[:1]]
    )
    lead_html = ""
    if candidati:
        lead, tag = candidati[0]
        lead_html = f"""
  <div class="lead"><span class="tag">{tag}</span>
    <span class="titolo">{_clean_title(lead.get("title", ""))}</span></div>"""

    meteo_html = "".join(_render_meteo_blocco(m) for m in meteo)
    return f"""
<div class="page page-front">
  <h1 class="masthead">NT REPORT</h1>
  <div class="data">{datetime.now().strftime("%A %d %B %Y")}</div>
  {lead_html}
  {_PRIVATO_INIZIO}
  {meteo_html}
  {_build_todo_checklist(todo)}
  {_PRIVATO_FINE}
</div>"""


def _render_item_li(it: dict) -> str:
    riassunto = it.get("riassunto", "")
    riassunto_html = f'<span class="riassunto">{riassunto}</span>' if riassunto else ""
    fonte = it.get("source_name", "")
    fonte_html = f'<span class="fonte">{fonte}</span>' if fonte else ""
    return (
        f'<li><a href="{it.get("url","")}">{_clean_title(it.get("title",""))}</a>'
        f'{riassunto_html}{fonte_html}</li>'
    )


def _frammento_settori(settori: dict, giustizia_amm: list) -> str:
    blocchi = []
    for label, items in settori.items():
        if not items:
            continue
        righe = "".join(_render_item_li(it) for it in items)
        blocchi.append(f'<div class="settore-blocco"><h3>{label}</h3><ul>{righe}</ul></div>')
    corpo = "".join(blocchi) or "<p><em>Nessuna novità nelle fonti professionali oggi.</em></p>"

    if giustizia_amm:
        righe = "".join(_render_item_li(it) for it in giustizia_amm)
        corpo += f'<div class="settore-blocco"><h3>Giustizia Amministrativa — Ufficio Studi</h3><ul>{righe}</ul></div>'

    return f'<h2 class="sezione">Professionale — Energia, Giochi, Concessioni, Tecnologia</h2>{corpo}'


def _frammento_interessi(interessi: list) -> str:
    if not interessi:
        corpo = "<p><em>Nessuna notizia particolarmente rilevante oggi su questi temi.</em></p>"
    else:
        by_tema: dict[str, list] = {}
        for it in interessi:
            by_tema.setdefault(it.get("tema") or "Altro", []).append(it)
        blocchi = []
        for tema, items in by_tema.items():
            righe = "".join(_render_item_li(it) for it in items)
            blocchi.append(f'<div class="settore-blocco"><h3>{tema.capitalize()}</h3><ul>{righe}</ul></div>')
        corpo = "".join(blocchi)

    return f'<h2 class="sezione">Interessi — Geopolitica, Difesa, Spazio, Scienza</h2>{corpo}'


def _frammento_attualita(attualita: list, ft_stampa: list) -> str:
    righe = "".join(_render_item_li(it) for it in attualita)
    corpo = f"<ul>{righe}</ul>" if attualita else "<p><em>Nessuna notizia trovata oggi.</em></p>"

    if ft_stampa:
        righe_ft = "".join(
            f'<li>{it["testo"]}' + (f' — <a href="{it["url"]}">link</a>' if it.get("url") else "") + "</li>"
            for it in ft_stampa
        )
        corpo += f'<div class="settore-blocco"><h3>Dal Financial Times e La Stampa (segnalati a mano)</h3><ul>{righe_ft}</ul></div>'

    return f'<h2 class="sezione">Attualità</h2>{corpo}'


def _frammento_linkedin(vetrina: list[dict]) -> str:
    """Vetrina di post/spunti raccolti nell'ultimo giro supervisionato su LinkedIn
    (richiesta di Niccolò, 2026-09-16). Non appare affatto se non c'è un giro abbastanza
    recente (vedi _load_linkedin_vetrina_recente) — niente sezione vuota o con contenuto
    stantio."""
    if not vetrina:
        return ""
    righe = "".join(
        "<li>"
        + it["testo"]
        + (f' — <span class="fonte">{it["autore"]}</span>' if it.get("autore") else "")
        + (f' <a href="{it["url"]}">link</a>' if it.get("url") else "")
        + "</li>"
        for it in vetrina
    )
    return f'<h2 class="sezione">Dal mio giro su LinkedIn</h2><ul>{righe}</ul>'


def _render_varie_block(label: str, dato: dict | None) -> str:
    if not dato or not dato.get("testo"):
        return ""
    html = f'<div class="varie-block"><span class="label">{label}:</span> {dato["testo"]}'
    if dato.get("link"):
        html += f' — <a href="{dato["link"]}">{dato["link"]}</a>'
    if dato.get("immagine"):
        html += f'<br><img class="varie-img" src="{dato["immagine"]}">'
    html += "</div>"
    return html


def _render_meteo_fascia(nome: str, dato: dict | None) -> str:
    if not dato:
        return f'<td class="meteo-fascia"><div class="meteo-fascia-nome">{nome}</div></td>'
    icona = _METEO_ICONE.get(dato.get("condizione", ""), "")
    temp = dato.get("temperatura", "")
    vento = ""
    if dato.get("vento_intensita") or dato.get("vento_direzione"):
        vento = f'<div class="meteo-vento">Vento: {dato.get("vento_intensita","")} {dato.get("vento_direzione","")}</div>'
    return f"""<td class="meteo-fascia">
      <div class="meteo-fascia-nome">{nome}</div>
      <div class="meteo-icona">{icona}</div>
      <div class="meteo-temp">{temp}</div>
      {vento}
    </td>"""


def _render_meteo_blocco(m: dict) -> str:
    intestazione = f'<div class="label">Meteo domani — {m["citta"]}</div>'
    riga_sole = (
        f'<div class="meteo-sole">Alba {m.get("alba","?")} · Tramonto {m.get("tramonto","?")}'
        + (f' · Umidità {m["umidita"]}' if m.get("umidita") else "")
        + "</div>"
    )
    celle = "".join(
        _render_meteo_fascia(nome, m.get(chiave))
        for nome, chiave in (("Mattina", "mattina"), ("Pomeriggio", "pomeriggio"),
                              ("Sera", "sera"), ("Notte", "notte"))
    )
    return f"""
    <div class="varie-block meteo-blocco">
      {intestazione}
      {riga_sole}
      <table class="meteo-tabella"><tr>{celle}</tr></table>
    </div>"""


def _build_ultima_pagina(varie: dict, varie_viaggio: dict | None) -> str:
    """Ultima pagina del giornale: spunti di Varie e viaggio tutti insieme (richiesta di
    Niccolò, 2026-09-16 — prima erano sparsi tra il flusso continuo e una pagina a parte).
    Il meteo è passato in prima pagina e il cruciverba è stato tolto (richiesta del
    2026-10-08). Blocco singolo, niente colonne: dentro un contenitore multi-colonna
    paginato su più pagine fisiche "column-span: all" manda in crash WeasyPrint in modo
    intermittente (IndexError interno in skip_first_whitespace — bug del motore, non del
    nostro CSS; riprodotto in produzione il 2026-09-15 e di nuovo il 2026-09-16)."""
    # "A tavola" per ultimo tra gli spunti (richiesta di Niccolò, 2026-09-15).
    labels = {
        "mostra": "Da vedere", "attivita_bambini": "Con i bambini",
        "libro": "Da leggere", "disco": "Da ascoltare", "film_serie": "Da guardare",
        "piatto_vino": "A tavola",
    }
    varie_html = "".join(_render_varie_block(label, varie.get(key)) for key, label in labels.items())

    viaggio_html = ""
    if varie_viaggio:
        labels_viaggio = {"fare": "Cosa fare", "vedere": "Cosa vedere", "mangiare": "Dove mangiare"}
        blocchi_viaggio = "".join(
            _render_varie_block(label, varie_viaggio.get(key)) for key, label in labels_viaggio.items()
        )
        viaggio_html = f'<h3>In viaggio a {varie_viaggio["citta"]}</h3>{blocchi_viaggio}'

    return (
        '<div class="page"><h2 class="sezione">Varie</h2>'
        + viaggio_html
        + varie_html
        + "</div>"
    )


def _build_corpo_continuo(settori, giustizia_amm, interessi, attualita, ft_stampa,
                           linkedin_vetrina) -> str:
    contenuto = (
        _frammento_settori(settori, giustizia_amm)
        + _frammento_interessi(interessi)
        + _frammento_attualita(attualita, ft_stampa)
        + _frammento_linkedin(linkedin_vetrina)
    )
    return f'<div class="page"><div class="corpo-continuo">{contenuto}</div></div>'


def _build_html(settori, interessi, attualita, giustizia_amm, ft_stampa, varie,
                 todo, meteo, varie_viaggio, linkedin_vetrina) -> str:
    return f"""<!doctype html>
<html lang="it">
<head>
<meta charset="utf-8">
<style>{_CSS}</style>
</head>
<body>
  {_build_front_page(settori, interessi, attualita, meteo, todo)}
  {_build_corpo_continuo(settori, giustizia_amm, interessi, attualita, ft_stampa, linkedin_vetrina)}
  {_build_ultima_pagina(varie, varie_viaggio)}
</body>
</html>"""


# ── Immagine del giorno (solo versione pubblica) ─────────────────────────────────────
# Nella copia per nt-report.com meteo e adempimenti spariscono e al loro posto va
# un'immagine (richiesta di Niccolò, 2026-10-08): a giorni alterni un quadro famoso di
# pubblico dominio e un'immagine generata con AI, sempre con una didascalia che dice cos'è.
# Se la fonte del giorno non risponde si prova la successiva; se falliscono tutte la prima
# pagina pubblica esce con il solo titolo di apertura.
# Quadri: Wikidata + Wikimedia Commons; riserva Cleveland Museum of Art (CC0). Art
# Institute of Chicago e Met sono stati scartati il 2026-10-08: le loro immagini
# rispondono 403 alle richieste non da browser.
# Wikimedia chiede uno User-Agent descrittivo con un contatto, altrimenti risponde 403.
_HTTP_UA = {"User-Agent": "NT-Report/1.0 (rassegna quotidiana; https://nt-report.com)"}


def _img_data_uri(contenuto: bytes, taglia_basso: float = 0.0) -> str:
    """Ridimensiona per la stampa e incorpora l'immagine nell'HTML: WeasyPrint non deve
    andare in rete mentre impagina. taglia_basso toglie una fascia in fondo (serve per il
    logo che Pollinations stampa nell'angolo in basso a destra)."""
    import base64
    import io
    from PIL import Image

    img = Image.open(io.BytesIO(contenuto)).convert("RGB")
    if taglia_basso:
        img = img.crop((0, 0, img.width, int(img.height * (1 - taglia_basso))))
    img.thumbnail((1400, 1400))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()


def _quadri_famosi() -> list[dict]:
    """Elenco di dipinti famosi di pubblico dominio da Wikidata: "famoso" vuol dire con
    almeno 20 voci Wikipedia/Wikimedia collegate, "pubblico dominio" vuol dire autore
    morto da oltre 100 anni. La query va su QLever (Università di Friburgo, stessi dati di
    Wikidata) e non sul servizio SPARQL ufficiale: il 2026-10-08 quello impiegava 50
    secondi o andava in timeout anche su query ridotte, QLever risponde in un secondo e
    mezzo. L'elenco va comunque in cache su disco e si aggiorna una volta al mese; se
    l'aggiornamento fallisce si continua con la cache vecchia."""
    import json
    import httpx

    cache = BASE_DIR / "data" / "quadri_famosi.json"
    if cache.exists() and (datetime.now().timestamp() - cache.stat().st_mtime) < 30 * 86400:
        return json.loads(cache.read_text(encoding="utf-8"))

    anno_limite = datetime.now().year - 100
    query = f"""PREFIX wd: <http://www.wikidata.org/entity/>
PREFIX wdt: <http://www.wikidata.org/prop/direct/>
PREFIX schema: <http://schema.org/>
PREFIX rdfs: <http://www.w3.org/2000/01/rdf-schema#>
SELECT ?item (SAMPLE(?t) AS ?titolo) (SAMPLE(?au) AS ?autore) (MIN(?an) AS ?anno)
       (SAMPLE(?im) AS ?img) WHERE {{
  ?item wdt:P31 wd:Q3305213; wdt:P18 ?im; wdt:P170 ?a.
  ?a wdt:P570 ?morte. FILTER(YEAR(?morte) < {anno_limite})
  ?voce schema:about ?item.
  ?item rdfs:label ?t. FILTER(LANG(?t) = "it")
  ?a rdfs:label ?au. FILTER(LANG(?au) = "it")
  OPTIONAL {{ ?item wdt:P571 ?d. BIND(YEAR(?d) AS ?an) }}
}} GROUP BY ?item HAVING (COUNT(DISTINCT ?voce) >= 20)"""
    try:
        r = httpx.get("https://qlever.dev/api/wikidata", params={"query": query},
                      headers={**_HTTP_UA, "Accept": "application/sparql-results+json"}, timeout=120)
        r.raise_for_status()
        elenco = sorted(
            ({"titolo": b["titolo"]["value"], "autore": b["autore"]["value"],
              "anno": b.get("anno", {}).get("value", ""), "img": b["img"]["value"]}
             for b in r.json()["results"]["bindings"]),
            key=lambda q: q["img"],
        )
        if not elenco:
            raise ValueError("risposta vuota")
        cache.parent.mkdir(exist_ok=True)
        cache.write_text(json.dumps(elenco, ensure_ascii=False), encoding="utf-8")
        return elenco
    except Exception as e:
        logger.warning(f"Rassegna: aggiornamento elenco quadri famosi da Wikidata fallito: {e}")
        return json.loads(cache.read_text(encoding="utf-8")) if cache.exists() else []


def _quadro_famoso_del_giorno() -> dict | None:
    import httpx

    quadri = _quadri_famosi()
    if not quadri:
        return None
    # Ordine mescolato una volta con seme fisso e scorso un passo per edizione: nessun
    # quadro si ripete prima di aver visto tutti gli altri. Il passo è ordinal // 2
    # perché le opere escono un giorno sì e uno no (vedi _blocco_arte_html).
    random.Random(0).shuffle(quadri)
    q = quadri[(datetime.now().toordinal() // 2) % len(quadri)]
    try:
        img = httpx.get(q["img"], params={"width": 1400}, headers=_HTTP_UA,
                        timeout=60, follow_redirects=True)
        img.raise_for_status()
    except Exception as e:
        logger.warning(f"Rassegna: immagine di '{q['titolo']}' non scaricabile: {e}")
        return None
    anno = f", {q['anno']}" if q["anno"] else ""
    return {"src": _img_data_uri(img.content),
            "didascalia": f"Opera del giorno — {q['autore']}, {q['titolo']}{anno}. "
                          f"Pubblico dominio, via Wikimedia Commons."}


def _opera_cleveland() -> dict | None:
    """Riserva per quando Wikidata/Commons non rispondono: un dipinto qualsiasi (non
    necessariamente famoso) dalla collezione Open Access del Cleveland Museum of Art."""
    import httpx

    base = "https://openaccess-api.clevelandart.org/api/artworks/"
    filtri = {"cc0": 1, "has_image": 1, "type": "Painting"}
    try:
        totale = httpx.get(base, params={**filtri, "limit": 1}, timeout=30).json()["info"]["total"]
        # Scelta deterministica per data: rigenerare la stessa edizione dà la stessa opera.
        indice = random.Random(datetime.now().strftime("%Y-%m-%d")).randrange(totale)
        opera = httpx.get(base, params={**filtri, "limit": 1, "skip": indice},
                          timeout=30).json()["data"][0]
        img = httpx.get(opera["images"]["web"]["url"], headers=_HTTP_UA, timeout=60)
        img.raise_for_status()
    except Exception as e:
        logger.warning(f"Rassegna: opera del giorno non disponibile: {e}")
        return None
    autore = (opera.get("creators") or [{}])[0].get("description") or "Autore ignoto"
    titolo = opera["title"] if len(opera["title"]) <= 80 else opera["title"][:80].rsplit(" ", 1)[0] + "…"
    didascalia = (f"Opera del giorno — {autore}, {titolo}, {opera.get('creation_date', '')}. "
                  f"The Cleveland Museum of Art, pubblico dominio (CC0).")
    return {"src": _img_data_uri(img.content), "didascalia": didascalia}


def _immagine_ai_del_giorno() -> dict | None:
    """Claude sceglie soggetto e stile e scrive la didascalia in italiano; l'immagine la
    genera Pollinations (servizio gratuito, nessuna chiave: il prompt va nell'URL)."""
    import urllib.parse
    import httpx
    from bot import call_claude

    prompt = """\
Inventa il soggetto di un'immagine artistica da mettere in prima pagina su un giornale
quotidiano, al posto di una fotografia. Soggetto e stile ogni volta diversi: paesaggi,
nature morte, architetture, scene marine, astrazioni, città immaginarie; stili pittorici,
incisione, acquerello, fotografia d'autore. Niente persone riconoscibili, niente testo
nell'immagine, niente marchi.

Rispondi SOLO con un oggetto JSON:
{"prompt_en": "prompt dettagliato in inglese per il generatore di immagini (max 60 parole)",
 "didascalia": "una frase in italiano che descrive soggetto e stile (max 20 parole)"}"""
    try:
        msg = call_claude(model="claude-haiku-4-5", max_tokens=300,
                          messages=[{"role": "user", "content": prompt}])
        idea = _estrai_json(msg.content[0].text)
        url = ("https://image.pollinations.ai/prompt/" + urllib.parse.quote(idea["prompt_en"])
               + f"?width=1200&height=800&nologo=true&seed={datetime.now().toordinal()}")
        img = httpx.get(url, headers=_HTTP_UA, timeout=120, follow_redirects=True)
        img.raise_for_status()
        if not img.headers.get("content-type", "").startswith("image/"):
            raise ValueError(f"risposta non è un'immagine: {img.headers.get('content-type')}")
    except Exception as e:
        logger.warning(f"Rassegna: immagine AI non disponibile: {e}")
        return None
    return {"src": _img_data_uri(img.content, taglia_basso=0.07),
            "didascalia": f"Immagine generata con AI — {idea['didascalia']}"}


def _blocco_arte_html() -> str:
    fonti = [_quadro_famoso_del_giorno, _immagine_ai_del_giorno]
    if datetime.now().toordinal() % 2:
        fonti.reverse()
    fonti.append(_opera_cleveland)
    for fonte in fonti:
        arte = fonte()
        if arte:
            return (f'<div class="arte"><img src="{arte["src"]}">'
                    f'<div class="didascalia">{arte["didascalia"]}</div></div>')
    return ""


# ── Orchestrazione ───────────────────────────────────────────────────────────────

def _versione_pubblica_pdf(pdf_path: Path) -> bytes:
    """Restituisce i bytes di una versione pubblicabile su nt-report.com, rigenerata
    dall'HTML sorgente (salvato accanto al PDF da genera_rassegna_pdf) rimuovendo il
    blocco privato della prima pagina (meteo + adempimenti, tra _PRIVATO_INIZIO e
    _PRIVATO_FINE) PRIMA di renderizzare il PDF, e mettendo al suo posto l'immagine del
    giorno (_blocco_arte_html). Non si basa
    sul conteggio delle pagine: se rimuovessimo solo "la prima pagina" del PDF finale,
    un'edizione con più notizie del solito potrebbe far sconfinare l'indice/ToDo sulla
    pagina 2, lasciando la ToDo list visibile pubblicamente (bug corretto il 22/09/2026,
    vedi feedback_privacy_rassegna in memoria). La copia stampata/su Drive (pdf_path)
    resta invariata: questa funzione non la tocca."""
    import re
    from weasyprint import HTML

    html_path = pdf_path.with_suffix(".html")
    if html_path.exists():
        html = html_path.read_text(encoding="utf-8")
        if _PRIVATO_INIZIO in html:
            inizio = html.index(_PRIVATO_INIZIO)
            fine = html.index(_PRIVATO_FINE) + len(_PRIVATO_FINE)
            html_pubblico = html[:inizio] + _blocco_arte_html() + html[fine:]
        else:
            # Edizioni fino al 2026-10-08, prima dei marcatori: ToDo in un div senza
            # div annidati, quindi il regex non annidato basta.
            html_pubblico = re.sub(r'<div class="todo">.*?</div>', "", html, flags=re.DOTALL)
        if (_PRIVATO_INIZIO in html_pubblico or '<div class="todo">' in html_pubblico
                or 'class="varie-block meteo-blocco"' in html_pubblico):
            # Rimozione incompleta (es. markup cambiato): meglio fallire rumorosamente che
            # pubblicare per errore un HTML ancora privato.
            raise ValueError("Rimozione del blocco privato fallita — controllare _build_front_page")
        return HTML(string=html_pubblico).write_pdf()

    # Fallback per PDF generati prima che genera_rassegna_pdf salvasse anche l'HTML
    # sorgente (edizioni del 14-21/09/2026): cerca ed esclude qualsiasi pagina il cui
    # testo contenga il marcatore "DA FARE", invece di assumere sia sempre la pagina 1.
    return _versione_pubblica_pdf_fallback_pypdf(pdf_path)


def _versione_pubblica_pdf_fallback_pypdf(pdf_path: Path) -> bytes:
    import io
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(str(pdf_path))
    writer = PdfWriter()
    for page in reader.pages:
        testo = (page.extract_text() or "").upper()
        if "DA FARE" in testo:
            continue
        writer.add_page(page)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _pubblica_rassegna_sul_sito(pdf_path: Path) -> bool:
    """Carica la versione pubblica (senza meteo e adempimenti) su nt-report.com via l'endpoint
    admin del sito. Richiede ADMIN_TOKEN nell'ambiente, condiviso con il servizio Render
    nt-report-api. Non blocca il resto del job se fallisce."""
    import os
    import httpx

    token = os.environ.get("ADMIN_TOKEN", "")
    if not token:
        logger.warning("Rassegna: ADMIN_TOKEN non configurato — salto pubblicazione sul sito")
        return False

    data_str = pdf_path.stem
    try:
        pdf_bytes = _versione_pubblica_pdf(pdf_path)
        resp = httpx.post(
            "https://nt-report-api-kojk.onrender.com/api/admin/rassegne",
            headers={"X-Admin-Token": token},
            data={"data": data_str},
            files={"file": (f"{data_str}.pdf", pdf_bytes, "application/pdf")},
            timeout=30,
        )
        resp.raise_for_status()
        return True
    except Exception as e:
        logger.warning(f"Rassegna: pubblicazione sul sito fallita: {e}")
        return False


def _pubblica_varie_sul_sito(data: str | None = None, tutte: bool = False) -> bool:
    """Invia le varie del giorno (o, con tutte=True, l'intero varie_storico — recupero
    una tantum) alla webapp nt-report.com/varie.html, dove Niccolò le spunta, le vota da
    1 a 3 stelle e le commenta dal cellulare (richiesta del 2026-09-29). Legge da
    varie_storico, non dai dict in memoria, così giorno corrente e recupero storico
    passano dallo stesso percorso. Il sito fa upsert senza toccare spunte/stelle/commenti,
    quindi rilanciarla è innocuo. Non blocca il resto del job se fallisce."""
    import os
    import httpx

    token = os.environ.get("ADMIN_TOKEN", "")
    if not token:
        logger.warning("Varie: ADMIN_TOKEN non configurato — salto pubblicazione sul sito")
        return False

    con = _varie_storico_connect()
    if tutte:
        righe = con.execute(
            "SELECT data, chiave, testo, link, immagine, citta FROM varie_storico"
        ).fetchall()
    else:
        data = data or datetime.now().strftime("%Y-%m-%d")
        righe = con.execute(
            "SELECT data, chiave, testo, link, immagine, citta FROM varie_storico WHERE data = ?",
            (data,),
        ).fetchall()
    con.close()
    if not righe:
        logger.warning("Varie: nessuna voce da pubblicare")
        return False

    voci = [
        {"data": d, "categoria": k, "testo": t, "link": l, "immagine": i, "citta": c}
        for d, k, t, l, i, c in righe
    ]
    try:
        resp = httpx.post(
            "https://nt-report-api-kojk.onrender.com/api/admin/varie",
            headers={"X-Admin-Token": token},
            json={"voci": voci},
            timeout=30,
        )
        resp.raise_for_status()
        return True
    except Exception as e:
        logger.warning(f"Varie: pubblicazione sul sito fallita: {e}")
        return False


async def run_rassegna_job(context) -> None:
    """Job giornaliero: genera il PDF, lo carica su Drive (per il print agent locale) e lo
    invia come documento Telegram all'owner. Solo per l'owner, non blocca in caso di errori
    parziali (es. upload Drive fallito) — avvisa e prosegue dove possibile."""
    from monitor import _upload_rassegna_to_drive, _upload_rassegna_settimana_to_drive

    chat_id = context.bot_data.get("owner_chat_id")
    if not chat_id:
        logger.warning("Rassegna: owner_chat_id non trovato — manda /start al bot")
        return

    try:
        pdf_path = await genera_rassegna_pdf()
    except Exception as e:
        logger.error(f"Rassegna: generazione PDF fallita: {e}", exc_info=True)
        await context.bot.send_message(chat_id=chat_id, text=f"⚠️ Rassegna: generazione PDF fallita: {str(e)[:300]}")
        return

    import asyncio
    loop = asyncio.get_event_loop()
    uploaded = await loop.run_in_executor(None, _upload_rassegna_to_drive, pdf_path)
    if uploaded:
        await context.bot.send_message(chat_id=chat_id, text="📂 Rassegna caricata su Drive per la stampa.")
    else:
        await context.bot.send_message(
            chat_id=chat_id,
            text="⚠️ Upload su Drive saltato o fallito (vedi log) — la stampa automatica non partirà.",
        )

    settimana_ok = await loop.run_in_executor(None, _upload_rassegna_settimana_to_drive, pdf_path)
    if not settimana_ok:
        await context.bot.send_message(
            chat_id=chat_id,
            text="⚠️ Archivio settimanale su Drive saltato o fallito (vedi log).",
        )

    pubblicata = await loop.run_in_executor(None, _pubblica_rassegna_sul_sito, pdf_path)
    if pubblicata:
        await context.bot.send_message(chat_id=chat_id, text="🌐 Rassegna pubblicata su nt-report.com.")
    else:
        await context.bot.send_message(
            chat_id=chat_id,
            text="⚠️ Pubblicazione su nt-report.com saltata o fallita (vedi log).",
        )

    varie_ok = await loop.run_in_executor(None, _pubblica_varie_sul_sito)
    if not varie_ok:
        await context.bot.send_message(
            chat_id=chat_id,
            text="⚠️ Pubblicazione delle varie su nt-report.com/varie.html saltata o fallita (vedi log).",
        )

    with open(pdf_path, "rb") as f:
        await context.bot.send_document(chat_id=chat_id, document=f, filename=pdf_path.name)


async def genera_rassegna_html() -> str:
    """Esegue tutti i fetch/generazioni e ritorna l'HTML completo (usato sia dal path di
    produzione — che poi lo converte in PDF con WeasyPrint — sia da test locali)."""
    settori = await _fetch_settori_professionali()
    interessi = await _fetch_interessi()
    attualita = await _fetch_attualita()
    giustizia_amm = await _fetch_giustizia_amministrativa()
    ft_stampa = _get_ft_stampa_queue()
    todo = _fetch_todo_list()
    citta_viaggio = _rileva_citta_viaggio(todo)
    todo = sorted(_unisci_e_accorcia(todo), key=lambda v: v["data"])
    varie = _genera_varie()
    varie_viaggio = _fetch_varie_viaggio(citta_viaggio)
    meteo = await _fetch_meteo(citta_viaggio)
    linkedin_vetrina = _load_linkedin_vetrina_recente()

    # Il PDF va stampato: ogni voce deve avere un riassunto leggibile, non solo un link.
    settori_items = [it for items in settori.values() for it in items]
    await _arricchisci_con_riassunti(settori_items, interessi, attualita, giustizia_amm)

    return _build_html(settori, interessi, attualita, giustizia_amm, ft_stampa, varie,
                        todo, meteo, varie_viaggio, linkedin_vetrina)


async def genera_rassegna_pdf() -> Path:
    """Genera il PDF del giorno e lo salva in rassegna/YYYY-MM-DD.pdf. Non committato su
    git (artefatto binario giornaliero, escluso via .gitignore). Salva anche l'HTML
    sorgente accanto al PDF (stesso nome, .html): serve a _versione_pubblica_pdf per
    ricostruire una copia senza meteo e adempimenti privati, senza dover contare le pagine."""
    from weasyprint import HTML

    html = await genera_rassegna_html()

    RASSEGNA_DIR.mkdir(exist_ok=True)
    date_str = datetime.now().strftime('%Y-%m-%d')
    out_path = RASSEGNA_DIR / f"{date_str}.pdf"
    HTML(string=html).write_pdf(str(out_path))
    (RASSEGNA_DIR / f"{date_str}.html").write_text(html, encoding="utf-8")
    return out_path
